import os
import re
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import shutil

# ============================================================
# CONFIGURATION
# ============================================================
SKIP_CATS = ["DEFOPENS", "DEFOBCS", "DEFROBCS", "DEFLOPEN", "DEFLOBC"]
TABLE_COLS = ["Branch (Intake)", "Gender", "OPEN", "SC", "ST", "VJ-DT",
              "NT-1 (B)", "NT-2 (C)", "NT-3 (D)", "OBC", "SEBC", "EWS", "TFWS", "AI (JEE)"]
CAT_ORDER = ["OPEN", "SC", "ST", "VJ-DT", "NT-1 (B)", "NT-2 (C)", "NT-3 (D)", "OBC", "SEBC", "EWS", "TFWS"]

# ============================================================
# HELPER FUNCTIONS
# ============================================================
def get_file_path(prompt, allowed_exts):
    while True:
        path = input(prompt).strip().strip('"').strip("'")
        if os.path.isfile(path):
            _, ext = os.path.splitext(path)
            if ext.lower() in allowed_exts: return path
        print("  ✘ File not found. Please try again.")

def find_col(df, keywords):
    """Finds a column that contains all the specified keywords."""
    for col in df.columns:
        col_str = str(col).upper()
        if all(kw.upper() in col_str for kw in keywords):
            return col
    return None

def parse_int(val):
    if val is None: return None
    if isinstance(val, (int, float)):
        if pd.isna(val): return None
        return int(val)
    
    s = str(val).strip()
    if not s: return None
    
    match = re.search(r'-?\d+', s)
    if match:
        try: return int(match.group())
        except: return None
    return None

def get_cutoff_rank(r1, r2, r3, r4):
    vals = [v for v in [parse_int(r1), parse_int(r2), parse_int(r3), parse_int(r4)] if v is not None]
    if not vals: return "-----"
    return max(vals)

def set_cell_shading(cell, color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>'))

def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    b = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                  f'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                  f'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                  f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                  f'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                  f'</w:tcBorders>')
    tcPr.append(b)

def make_run(p, text, bold=False, size=8):
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Arial'
    run.font._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    return run

def add_page_break(doc):
    p = doc.add_paragraph()
    p._element.clear()
    p._element.append(parse_xml(f'<w:r {nsdecls("w")}><w:br w:type="page"/></w:r>'))

def classify_and_route(cat):
    c = str(cat).strip().upper()
    if not c or c in SKIP_CATS: return None
    
    is_male = c.startswith(("GOPEN", "GSC", "GST", "GVJ", "GNT1", "GNT2", "GNT3", "GOBC", "GSEBC"))
    gender = "Male" if is_male else "Female"
    
    base = None
    if "EWS" in c: base = "EWS"
    elif "TFWS" in c: base = "TFWS"
    elif "OPEN" in c: base = "OPEN"
    elif c.startswith(("GSC", "LSC")): base = "SC"
    elif c.startswith(("GST", "LST")): base = "ST"
    elif "VJ" in c: base = "VJ-DT"
    elif "NT1" in c: base = "NT-1 (B)"
    elif "NT2" in c: base = "NT-2 (C)"
    elif "NT3" in c: base = "NT-3 (D)"
    elif "OBC" in c: base = "OBC"
    elif "SEBC" in c: base = "SEBC"
    if not base: return None
    
    if base in ("EWS", "TFWS"): return (gender, base, "BOTH")
    
    if c.endswith("S"): return (gender, base, "STATE")
    if c.endswith("O"): return (gender, base, "OHU")
    return (gender, base, "HU")

# ============================================================
# EXCEL PARSER
# ============================================================
def process_excel(df):
    cat_col = find_col(df, ["CATEGORY"])
    
    # ONLY look for Branch Name / Course Name columns
    branch_name_col = find_col(df, ["BRANCH", "NAME"]) or find_col(df, ["COURSE", "NAME"])
    
    # Fallback if only a general "Branch" column exists (Ensure it's not the Code column)
    if not branch_name_col:
        for col in df.columns:
            col_str = str(col).upper()
            if ("BRANCH" in col_str or "COURSE" in col_str) and "CODE" not in col_str:
                branch_name_col = col
                break
                
    code_col = find_col(df, ["INSTITUTE", "CODE"])
    inst_col = find_col(df, ["INSTITUTE"])
    
    r1_col = find_col(df, ["R1", "RANK"])
    r2_col = find_col(df, ["R2", "RANK"])
    r3_col = find_col(df, ["R3", "RANK"])
    r4_col = find_col(df, ["R4", "RANK"])
    r_cols = [r1_col, r2_col, r3_col, r4_col]

    if not cat_col or not r1_col: return None, None

    college_name = ""
    if code_col and inst_col:
        cv = str(df[code_col].dropna().iloc[0]).strip() if not df[code_col].dropna().empty else ""
        iv = str(df[inst_col].dropna().iloc[0]).strip() if not df[inst_col].dropna().empty else ""
        if cv and iv and iv.lower() != "nan": college_name = f"{cv}-{iv}"

    branches = {}
    cur_branch_name = ""
    
    for _, row in df.iterrows():
        # Persist Branch Name across rows if they are merged/blank in Excel
        if branch_name_col and pd.notna(row.get(branch_name_col)):
            val = str(row[branch_name_col]).strip()
            if val and val.lower() != "nan":
                cur_branch_name = val
                
        if not cur_branch_name: continue
        
        cat = row.get(cat_col)
        if pd.isna(cat): continue
        
        res = classify_and_route(cat)
        if not res: continue
        
        gender, base, dest = res
        
        ranks = [row.get(rc) if rc else None for rc in r_cols]
        rank = get_cutoff_rank(*ranks)
        
        if rank == "-----": continue
        
        if cur_branch_name not in branches:
            branches[cur_branch_name] = {"HU_M": {}, "HU_F": {}, "OHU_M": {}, "OHU_F": {}, "STATE_M": {}, "STATE_F": {}}
        
        b = branches[cur_branch_name]
        if dest in ("HU", "BOTH"):
            k = "HU_M" if gender == "Male" else "HU_F"
            if base not in b[k]: b[k][base] = rank
        if dest in ("OHU", "BOTH"):
            k = "OHU_M" if gender == "Male" else "OHU_F"
            if base not in b[k]: b[k][base] = rank
        if dest == "STATE":
            k = "STATE_M" if gender == "Male" else "STATE_F"
            if base not in b[k]: b[k][base] = rank
            
    return college_name, branches

# ============================================================
# WORD DOCUMENT BUILDERS
# ============================================================
def create_info_box(doc, name):
    tbl = doc.add_table(rows=3, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    texts = [
        [("College Code & Name: ", False), (name, False)],
        [("NBA Grading: ", True), ("", False)],
        [("Total Seats: ", True), ("", False)],
        [("Available Cut-off: ", True), ("", False)]
    ]
    widths = [Inches(4.0), Inches(2.0), Inches(1.5), Inches(2.5)]
    
    for i, parts in enumerate(texts):
        cell = tbl.cell(0, i)
        cell.width = widths[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
        for txt, b in parts: make_run(p, txt, bold=b)
        set_cell_shading(cell, "F2F2F2"); set_cell_border(cell)

    merged = tbl.cell(1, 0).merge(tbl.cell(2, 3))
    merged.text = ""
    for p in merged.paragraphs: p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    set_cell_shading(merged, "F2F2F2"); set_cell_border(merged)


def create_rank_table(doc, title, branches, mode):
    m_key, f_key = ("HU_M", "HU_F") if mode == "HU" else ("OHU_M", "OHU_F") if mode == "OHU" else ("STATE_M", "STATE_F")
    
    data_rows = []
    for b_name, b_data in branches.items():
        m_data, f_data = b_data[m_key], b_data[f_key]
        has_m = any(m_data.get(c) for c in CAT_ORDER)
        has_f = any(f_data.get(c) for c in CAT_ORDER)
        
        if has_m:
            row = [b_name, "Male"] + [m_data.get(c, "-----") for c in CAT_ORDER] + [""]
            data_rows.append(row)
        if has_f:
            row = ["", "Female"] + [f_data.get(c, "-----") for c in CAT_ORDER] + [""]
            data_rows.append(row)
            
    if not data_rows: return

    tbl = doc.add_table(rows=len(data_rows)+1, cols=14)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    w = [2.7, 0.7, 0.72, 0.72, 0.65, 0.65, 0.8, 0.8, 0.8, 0.65, 0.72, 0.65, 0.65, 0.65]
    for r in tbl.rows:
        for i, wi in enumerate(w): r.cells[i].width = Inches(wi)

    tbl._tbl.addprevious(
        parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:jc w:val="left"/><w:spacing w:before="120" w:after="60"/></w:pPr>'
                  f'<w:r><w:rPr><w:b/><w:sz w:val="22"/><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:eastAsia="Arial"/></w:rPr>'
                  f'<w:t>{title}</w:t></w:r></w:p>'))

    for i, col in enumerate(TABLE_COLS):
        c = tbl.rows[0].cells[i]; c.text = ""
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
        make_run(p, col, bold=True)
        set_cell_shading(c, "4472C4"); set_cell_border(c)
        for run in p.runs:
            run._element.get_or_add_rPr().append(parse_xml(f'<w:color {nsdecls("w")} w:val="FFFFFF"/>'))

    for ri, rd in enumerate(data_rows):
        bg = "FFFFFF" if ri % 2 == 0 else "F2F7FC"
        for ci in range(14):
            c = tbl.rows[ri+1].cells[ci]; c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
            val = rd[ci] if ci < len(rd) else ""
            if ci == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                make_run(p, val, bold=True)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                make_run(p, val, bold=False)
            set_cell_shading(c, bg); set_cell_border(c)

# ============================================================
# FOLDER & FILE MANAGEMENT
# ============================================================
def save_document(doc, college_name):
    folder_name = "College_Data"
    
    if not os.path.exists(folder_name):
        os.makedirs(folder_name)
        print(f"  ✔ Created folder: {folder_name}/")
    
    safe_name = re.sub(r'[<>:"/\\|?*]', '', college_name).strip()
    if not safe_name:
        safe_name = "Unknown_College"
        
    base_file_name = f"{safe_name}.docx"
    file_path = os.path.join(folder_name, base_file_name)
    
    if os.path.exists(file_path):
        counter = 1
        while True:
            new_name = f"{safe_name}_{counter}.docx"
            new_path = os.path.join(folder_name, new_name)
            if not os.path.exists(new_path):
                shutil.move(file_path, new_path)
                print(f"  ℹ Existing file renamed to: {new_name}")
                break
            counter += 1
            
    doc.save(file_path)
    print(f"\n  ✔ Saved new file as: {file_path}")

# ============================================================
# MAIN
# ============================================================
def main():
    print("=" * 60)
    print("  COLLEGE CUTOFF TABLE GENERATOR")
    print("=" * 60)
    
    excel_path = get_file_path("\nEnter path to Excel file (.xlsx/.xls): ", [".xlsx", ".xls"])
    print(f"  ✔ Loaded: {os.path.basename(excel_path)}")
    
    print("\nWhich table do you want to generate?")
    print("  1. Home University (HU)")
    print("  2. Other Than Home University (OHU)")
    print("  3. State Level")
    print("  4. HU & OHU Only")
    print("  5. All (HU, OHU & State)")
    
    while True:
        ch = input("Enter choice (1/2/3/4/5): ").strip()
        if ch in ("1", "2", "3", "4", "5"): break
    mode = {"1": "HU", "2": "OHU", "3": "STATE", "4": "HU_OHU", "5": "ALL"}[ch]

    print("\n  Processing...")
    try:
        df = pd.read_excel(excel_path)
    except Exception as e:
        print(f"  ✘ Error: {e}"); return

    name, branches = process_excel(df)
    if not branches:
        print("  ✘ No valid data found."); return

    print(f"  ✔ College: {name}")
    print(f"  ✔ Branches: {len(branches)}")

    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.top_margin, s.bottom_margin = Inches(0.5), Inches(0.45)
        s.left_margin, s.right_margin = Inches(0.5), Inches(0.5)

    st = doc.styles['Normal']
    st.font.name = 'Arial'; st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(0)

    if mode == "HU":
        create_info_box(doc, name)
        create_rank_table(doc, "Home University", branches, "HU")
    elif mode == "OHU":
        create_info_box(doc, name)
        create_rank_table(doc, "Other Than Home University", branches, "OHU")
    elif mode == "STATE":
        create_info_box(doc, name)
        create_rank_table(doc, "State Level", branches, "STATE")
    elif mode == "HU_OHU":
        create_info_box(doc, name)
        create_rank_table(doc, "Home University", branches, "HU")
        add_page_break(doc)
        create_info_box(doc, name)
        create_rank_table(doc, "Other Than Home University", branches, "OHU")
    elif mode == "ALL":
        create_info_box(doc, name)
        create_rank_table(doc, "Home University", branches, "HU")
        add_page_break(doc)
        create_info_box(doc, name)
        create_rank_table(doc, "Other Than Home University", branches, "OHU")
        add_page_break(doc)
        create_info_box(doc, name)
        create_rank_table(doc, "State Level", branches, "STATE")

    save_document(doc, name)
    print("=" * 60)

if __name__ == "__main__":
    main()